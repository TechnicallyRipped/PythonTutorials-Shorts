
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList

wb = load_workbook('scores_.xlsx')
ws = wb['Sheet1']

chart = BarChart()
chart.title = "Score by Name"
chart.x_axis.title = "Name"
chart.y_axis.title = "Score"

data = Reference(ws,min_col=2,min_row=1,max_row=5)

cats = Reference(ws,min_col=1,min_row=2,max_row=5)

chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)

series = chart.series[0]
series.dLbls = DataLabelList()
series.dLbls.showVal = True

ws.add_chart(chart, "A6")

wb.save('scores_chart2.xlsx')
